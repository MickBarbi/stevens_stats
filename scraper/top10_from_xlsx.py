"""Convert the official "Stevens T&F All-Time Top Ten Lists" workbook into
../data/top10.json (the shape the site expects — see data/README.md).

    pip install openpyxl
    python scraper/top10_from_xlsx.py path/to/Stevens_TF_AllTime_Top_Ten_Lists.xlsx

Only the four current "{Mens,Womens} {Indoor,Outdoor} Top Ten List" sheets are
read (the workbook also has "Outdoor Men" / "- Old" / "- New Print" sheets that
are stale or formatting-only duplicates — do not use them). Each is a grid of
4-column blocks: Rank, Mark, Name, Date.

Ties in the source have a blank Rank cell (same rank as the row above); those
rows are kept, not treated as the end of the list.
"""

from __future__ import annotations

import datetime
import json
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "scraper" / "Stevens_TF_AllTime_Top_Ten_Lists.xlsx"

SHEETS = [
    ("Mens Outdoor Top Ten List", "m", "outdoor"),
    ("Womens Outdoor Top Ten List", "f", "outdoor"),
    ("Mens Indoor Top Ten List", "m", "indoor"),
    ("Womens Indoor Top Ten List", "f", "indoor"),
]

# spreadsheet event label (whitespace/commas stripped, lowercased) -> (event_id, canonical)
EVENTS = {
    "100mdash": (2, "100 Meters"), "200mdash": (3, "200 Meters"), "400mdash": (4, "400 Meters"),
    "60mdash": (1, "60 Meters"), "55mdash": (32, "55 Meters"), "300mdash": (30, "300 Meters"),
    "500mrun": (31, "500 Meters"), "600mrun": (5, "600 Meters"), "800mrun": (6, "800 Meters"),
    "1000mrun": (7, "1000 Meters"), "1500mrun": (8, "1500 Meters"), "milerun": (9, "Mile"),
    "3000mrun": (10, "3000 Meters"), "5000mrun": (11, "5000 Meters"), "10000mrun": (12, "10,000 Meters"),
    "110mhurdles": (15, "110 Hurdles"), "100mhurdles": (14, "100 Hurdles"), "400mhurdles": (16, "400 Hurdles"),
    "60mhurdles": (13, "60 Hurdles"), "55mhurdles": (33, "55 Hurdles"),
    "3000msteeplechase": (17, "3000 Steeplechase"),
    "highjump": (18, "High Jump"), "polevault": (19, "Pole Vault"), "longjump": (20, "Long Jump"),
    "triplejump": (21, "Triple Jump"), "shotput": (22, "Shot Put"), "discus": (23, "Discus"),
    "hammer": (24, "Hammer"), "weightthrow": (25, "Weight Throw"), "javelin": (26, "Javelin"),
    "pentathlon": (27, "Pentathlon"), "heptathlon": (28, "Heptathlon"), "decathlon": (29, "Decathlon"),
}
RELAYS = {
    "4x100mrelay": "4x100 Relay", "4x200mrelay": "4x200 Relay", "4x400mrelay": "4x400 Relay",
    "4x800mrelay": "4x800 Relay", "distancemedleyrelay": "Distance Medley Relay",
}


def key(s: str) -> str:
    return re.sub(r"[\s,]+", "", s.lower())


def map_event(raw: str):
    base = re.sub(r"^(men's|women's)\s+", "", raw.strip(), flags=re.I)
    k = key(base)
    if k in EVENTS:
        return (*EVENTS[k], False)
    if k in RELAYS:
        return (None, RELAYS[k], True)
    return None


# ---- name -> current athlete id ------------------------------------------------
ATHLETES = json.loads((ROOT / "data" / "athletes.json").read_text("utf-8"))
BY_LAST: dict[str, list] = {}
for a in ATHLETES:
    BY_LAST.setdefault(a["last_name"].lower(), []).append(a)

ALIAS = {
    "danny": "daniel", "dan": "daniel", "matt": "matthew", "matty": "matthew", "chris": "christopher",
    "nick": "nicholas", "mike": "michael", "mikey": "michael", "will": "william", "willie": "william",
    "tom": "thomas", "tommy": "thomas", "joe": "joseph", "joey": "joseph", "jake": "jacob",
    "ben": "benjamin", "sam": "samuel", "alex": "alexander", "izzy": "isabella", "cate": "catherine",
    "kate": "katherine", "liz": "elizabeth", "gabe": "gabriel", "andy": "andrew", "tony": "anthony",
    "charlie": "charles", "greg": "gregory", "steve": "steven", "jon": "jonathan", "dave": "david",
    "nate": "nathan", "zach": "zachary", "josh": "joshua",
}


def first_ok(fl: str, af: str, nick: str | None) -> bool:
    if not fl or not af:
        return False
    if fl == af or af.startswith(fl) or fl.startswith(af):
        return True
    if nick and fl == nick.lower():
        return True
    if ALIAS.get(fl) == af or ALIAS.get(af) == fl:
        return True
    return len(fl) >= 3 and len(af) >= 3 and fl[:3] == af[:3]


def match_id(name: str):
    parts = name.strip().split()
    if len(parts) < 2:
        return None
    first, last = parts[0], " ".join(parts[1:])
    cands = BY_LAST.get(last.lower())
    if not cands:
        cands = BY_LAST.get(parts[-1].lower(), [])
        first = " ".join(parts[:-1])
    fl = first.lower()
    hits = [a for a in cands if first_ok(fl, a["first_name"].lower(), a.get("nickname"))]
    return hits[0]["athlete_id"] if len(hits) == 1 else None


# ---- cell helpers ------------------------------------------------------------
def clean_mark(m) -> str | None:
    if m is None:
        return None
    s = str(m).split("(")[0].strip().rstrip("*cChH ").strip()
    return s or None


def iso(d) -> str | None:
    if isinstance(d, datetime.datetime):
        return d.strftime("%Y-%m-%d")
    if isinstance(d, str) and re.match(r"\d{4}-\d\d-\d\d", d):
        return d[:10]
    return None


def parse_block(rows, header_i: int, j: int, relay: bool):
    """Read up to ~15 entries under a header, carrying a blank rank forward."""
    entries = []
    last_rank = 0
    for k in range(header_i + 1, min(header_i + 20, len(rows))):
        r = rows[k] or ()
        rc = r[j] if j < len(r) else None
        mark = clean_mark(r[j + 1] if j + 1 < len(r) else None)
        nm = r[j + 2] if j + 2 < len(r) else None
        nm = str(nm).strip() if nm not in (None, "") else None

        if isinstance(rc, str) and rc.strip() == "Rank":
            break  # ran into the next block
        if mark is None and nm is None:
            break  # blank separator row = end of this list

        if isinstance(rc, (int, float)):
            last_rank = int(rc)
        rank = last_rank or (len(entries) + 1)  # blank cell = tie with row above

        if mark is None or nm is None:
            continue  # malformed row, but list continues

        date = iso(r[j + 3] if j + 3 < len(r) else None)
        link = r[j + 4] if j + 4 < len(r) else None
        e = {"rank": rank, "athlete_id": None, "mark": mark}
        if relay:
            e["names"] = [x.strip() for x in nm.split(",") if x.strip()]
        else:
            e["name"] = nm
            aid = match_id(nm)
            if aid:
                e["athlete_id"] = aid
        if date:
            e["date"] = date
        if isinstance(link, str) and link.startswith("http"):
            e["link"] = link
        entries.append(e)
        if len(entries) >= 15:
            break
    return entries


def main(argv: list[str]) -> int:
    xlsx = pathlib.Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        sys.exit(f"workbook not found: {xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out, seen, unknown = [], set(), set()

    for sheet, gender, season in SHEETS:
        if sheet not in wb.sheetnames:
            print(f"  note: sheet {sheet!r} missing, skipped")
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        for i, row in enumerate(rows):
            if not row or i == 0:
                continue
            hdr_cols = [j for j, c in enumerate(row) if isinstance(c, str) and c.strip() == "Rank"]
            if not hdr_cols:
                continue
            ev_row = rows[i - 1] or ()
            for j in hdr_cols:
                raw = ev_row[j] if j < len(ev_row) else None
                if not (isinstance(raw, str) and raw.strip()):
                    continue
                mapped = map_event(raw)
                if not mapped:
                    unknown.add(raw.strip())
                    continue
                eid, canon, relay = mapped
                if (canon, gender, season) in seen:
                    continue
                entries = parse_block(rows, i, j, relay)
                if not entries:
                    continue
                seen.add((canon, gender, season))
                out.append({
                    "event_id": eid, "event_name": canon, "gender": gender,
                    "season": season, "relay": relay, "entries": entries,
                })

    out.sort(key=lambda l: (l["season"], l["gender"], l["relay"], l["event_id"] or 999, l["event_name"]))
    (ROOT / "data" / "top10.json").write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", "utf-8")

    names = {a["athlete_id"]: f"{a['first_name']} {a['last_name']}" for a in ATHLETES}
    links = [(l, e) for l in out for e in l["entries"] if e.get("athlete_id")]
    distinct = sorted({names[e["athlete_id"]] for _, e in links})
    print(f"lists={len(out)}  entries={sum(len(l['entries']) for l in out)}  "
          f"entry-links={len(links)}  distinct current athletes={len(distinct)}")
    short = [l for l in out if not l["relay"] and len(l["entries"]) < 10]
    if short:
        print(f"note: {len(short)} individual lists still have <10 entries "
              "(may be genuine — not every event has 10 all-time marks)")
    if unknown:
        print("UNMAPPED event labels:", sorted(unknown))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
